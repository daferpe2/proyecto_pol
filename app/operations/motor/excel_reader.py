from io import BytesIO
from openpyxl import load_workbook


class MissingColumnError(Exception):
    pass


class CorruptedExcelError(Exception):
    pass


class EmptyExcelError(Exception):
    pass


class ExcelReader:
    def read(self, file: bytes, expected_columns: list[str]):
        # Intentar abrir el archivo Excel
        try:
            workbook = load_workbook(filename=BytesIO(file), data_only=True)
        except Exception:
            raise CorruptedExcelError("El archivo Excel está corrupto o no se puede leer")

        sheet = workbook.active

        # Leer cabeceras
        headers = [cell.value for cell in sheet[1]]

        # Validar columnas obligatorias
        for col in expected_columns:
            if col not in headers:
                raise MissingColumnError(f"Falta la columna obligatoria: {col}")

        # Leer filas
        rows = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            rows.append(row_dict)

        # Validar que existan filas
        if not rows:
            raise EmptyExcelError("El archivo Excel no contiene filas")

        return rows, len(rows)

