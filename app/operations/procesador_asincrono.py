import asyncio
import io
import os
import uuid
from datetime import datetime

from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

from db import engine
from models import (
    BuscadorJob,
    JobEstado,
    TablaBaseEspF,
)

SEM_MAX = 10
RESULTS_DIR = "results"


async def procesar_excel_job(job_id: int, excel_bytes: bytes):
    """
    Procesa un archivo Excel asociado a un job:
    - Lee filas del Excel.
    - Busca en TablaBaseEspF usando texto_busqueda.
    - Genera un Excel de resultados.
    - Actualiza el job con ruta del archivo y estado final.
    """

    loop = asyncio.get_event_loop()
    sem = asyncio.Semaphore(SEM_MAX)

    # -----------------------------
    # Helpers síncronos de BD
    # -----------------------------
    def sync_get_job():
        with Session(engine) as session:
            return session.get(BuscadorJob, job_id)

    def sync_set_estado(estado: JobEstado, error: str | None = None):
        with Session(engine) as session:
            job = session.get(BuscadorJob, job_id)
            if not job:
                return
            job.estado = estado
            job.error_detalle = error
            job.fecha_actualizacion = datetime.now()
            session.add(job)
            session.commit()

    def sync_set_ruta(ruta: str):
        with Session(engine) as session:
            job = session.get(BuscadorJob, job_id)
            if not job:
                return
            job.ruta_archivo_resultado = ruta
            job.fecha_actualizacion = datetime.now()
            session.add(job)
            session.commit()

    # -----------------------------
    # 1. Obtener job y marcar EN_PROCESO
    # -----------------------------
    job = await loop.run_in_executor(None, sync_get_job)
    if not job:
        return

    await loop.run_in_executor(None, lambda: sync_set_estado(JobEstado.EN_PROCESO))

    try:
        # -----------------------------
        # 2. Cargar Excel desde bytes
        # -----------------------------
        wb = load_workbook(filename=io.BytesIO(excel_bytes), data_only=True)
        ws = wb.active

        # Validar encabezados mínimos
        encabezados = [cell.value for cell in ws[1]]
        if "no_cedula" not in encabezados or "apellidos_nombre" not in encabezados:
            raise ValueError("El Excel debe contener columnas 'no_cedula' y 'apellidos_nombre'.")

        idx_cedula = encabezados.index("no_cedula")
        idx_nombre = encabezados.index("apellidos_nombre")

        filas = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            filas.append({
                "no_cedula": str(row[idx_cedula] or "").strip(),
                "apellidos_nombre": str(row[idx_nombre] or "").strip(),
            })

        resultados: list[dict] = []

        # -----------------------------
        # 3. Función síncrona de búsqueda
        # -----------------------------
        def sync_buscar_por_criterio(criterio: str):
            with Session(engine) as session:
                stmt = select(TablaBaseEspF).where(
                    TablaBaseEspF.texto_busqueda.ilike(f"%{criterio}%")
                )
                return session.exec(stmt).all()

        # -----------------------------
        # 4. Procesar cada fila con semáforo
        # -----------------------------
        async def procesar_fila(fila: dict):
            async with sem:
                no_cedula = fila["no_cedula"]
                apellidos_nombre = fila["apellidos_nombre"]
                criterio = f"{no_cedula} {apellidos_nombre}".strip()

                try:
                    rows = await loop.run_in_executor(
                        None, lambda: sync_buscar_por_criterio(criterio)
                    )

                    if not rows:
                        resultados.append({
                            "input_no_cedula": no_cedula,
                            "input_apellidos_nombre": apellidos_nombre,
                            "criterio_busqueda": criterio,
                            "estado": "SIN_RESULTADOS",
                            "detalle_error": "",
                            "resultados": [],
                        })
                        return

                    registros = []
                    inconsistencia = False

                    for r in rows:
                        if r.no_cedula != no_cedula or r.apellidos_nombre != apellidos_nombre:
                            inconsistencia = True

                        registros.append({
                            "no_cedula": r.no_cedula,
                            "apellidos_nombre": r.apellidos_nombre,
                            "observacion": r.observacion,
                        })

                    estado = "INCONSISTENCIA_DATOS" if inconsistencia else "OK"
                    detalle = (
                        "La cédula y el nombre no coinciden con los registros encontrados"
                        if inconsistencia else ""
                    )

                    resultados.append({
                        "input_no_cedula": no_cedula,
                        "input_apellidos_nombre": apellidos_nombre,
                        "criterio_busqueda": criterio,
                        "estado": estado,
                        "detalle_error": detalle,
                        "resultados": registros,
                    })

                except Exception as e:
                    resultados.append({
                        "input_no_cedula": no_cedula,
                        "input_apellidos_nombre": apellidos_nombre,
                        "criterio_busqueda": criterio,
                        "estado": "ERROR_BD",
                        "detalle_error": str(e),
                        "resultados": [],
                    })

        await asyncio.gather(*(procesar_fila(f) for f in filas))

        # -----------------------------
        # 5. Generar Excel final
        # -----------------------------
        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.title = "Resultados"

        out_ws.append([
            "input_no_cedula",
            "input_apellidos_nombre",
            "criterio_busqueda",
            "estado",
            "detalle_error",
            "resultado_no_cedula",
            "resultado_apellidos_nombre",
            "resultado_observacion",
        ])

        for r in resultados:
            if not r["resultados"]:
                out_ws.append([
                    r["input_no_cedula"],
                    r["input_apellidos_nombre"],
                    r["criterio_busqueda"],
                    r["estado"],
                    r["detalle_error"],
                    "",
                    "",
                    "",
                ])
            else:
                for reg in r["resultados"]:
                    out_ws.append([
                        r["input_no_cedula"],
                        r["input_apellidos_nombre"],
                        r["criterio_busqueda"],
                        r["estado"],
                        r["detalle_error"],
                        reg["no_cedula"],
                        reg["apellidos_nombre"],
                        reg["observacion"],
                    ])

        # -----------------------------
        # 6. Guardar archivo
        # -----------------------------
        if not os.path.exists(RESULTS_DIR):
            os.makedirs(RESULTS_DIR, exist_ok=True)

        filename = f"resultado_job_{job_id}_{uuid.uuid4().hex}.xlsx"
        ruta = os.path.join(RESULTS_DIR, filename)

        out_wb.save(ruta)

        # -----------------------------
        # 7. Actualizar job
        # -----------------------------
        await loop.run_in_executor(None, lambda: sync_set_ruta(ruta))
        await loop.run_in_executor(None, lambda: sync_set_estado(JobEstado.COMPLETADO))

    except Exception:
        await loop.run_in_executor(None, lambda: sync_set_estado(JobEstado.ERROR, str(e)))
